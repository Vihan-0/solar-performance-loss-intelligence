"""
Generate Apple-Inspired 16:9 Presentation for Academic Evaluation Round II
Project: Adaptive Solar Performance & Loss Intelligence
"""
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Directory setup
REPO_ROOT = Path(__file__).resolve().parent.parent
PRESENTATION_DIR = REPO_ROOT / "presentation"
FIGURES_DIR = REPO_ROOT / "reports" / "figures"
PRESENTATION_DIR.mkdir(parents=True, exist_ok=True)

# -----------------------------------------------------------------------------
# APPLE-INSPIRED DESIGN SYSTEM TOKENS
# -----------------------------------------------------------------------------
COLOR_BG = RGBColor(11, 17, 20)           # #0B1114 Deep Midnight
COLOR_SURFACE = RGBColor(20, 30, 36)      # #141E24 Dark Graphite Card Fill
COLOR_SURFACE_ALT = RGBColor(26, 38, 46)  # #1A262E Slightly lighter card
COLOR_BORDER = RGBColor(35, 51, 62)       # #23333E Subtle card border
COLOR_TEXT_PRIMARY = RGBColor(248, 250, 252) # #F8FAFC Pure white text
COLOR_TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8 Slate muted text
COLOR_AMBER = RGBColor(245, 158, 11)      # #F59E0B Solar Amber / Risk
COLOR_GOLD = RGBColor(251, 191, 36)       # #FBBF24 Solar Gold
COLOR_EMERALD = RGBColor(16, 185, 129)    # #10B981 Emerald / Success
COLOR_CYAN = RGBColor(56, 189, 248)       # #38BDF8 Sky Cyan / Tech
COLOR_ROSE = RGBColor(244, 63, 94)        # #F43F5E Coral / Alert

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"

def set_slide_background(slide):
    """Fills slide background with deep midnight color."""
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLOR_BG
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, category_text, title_text):
    """Creates a consistent Apple-style top category tag and slide title."""
    # Category / Pill
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.733), Inches(0.3))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.name = FONT_HEADING
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_AMBER

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.733), Inches(0.55))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_TEXT_PRIMARY

def add_card(slide, left, top, width, height, bg_color=COLOR_SURFACE, border_color=COLOR_BORDER):
    """Draws a rounded rectangular surface card with subtle border."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    return card

# -----------------------------------------------------------------------------
# PRESENTATION INITIALIZATION
# -----------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# =============================================================================
# SLIDE 1: TITLE & PROJECT VISION
# =============================================================================
s1 = prs.slides.add_slide(blank_layout)
set_slide_background(s1)

# Subtle decorative background element
glow = s1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.5), Inches(-1.5), Inches(6.0), Inches(6.0))
glow.fill.solid()
glow.fill.fore_color.rgb = RGBColor(18, 30, 36)
glow.line.fill.background()

# Category Tag
tag_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(0.4))
tf_tag = tag_box.text_frame
tf_tag.word_wrap = True
p_tag = tf_tag.paragraphs[0]
p_tag.text = "ACADEMIC EVALUATION ROUND II • MINI-PROJECT DEFENSE"
p_tag.font.name = FONT_HEADING
p_tag.font.size = Pt(11)
p_tag.font.bold = True
p_tag.font.color.rgb = COLOR_AMBER

# Main Title
title_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(11.3), Inches(1.6))
tf_main = title_box.text_frame
tf_main.word_wrap = True
p_main = tf_main.paragraphs[0]
p_main.text = "Adaptive Solar Performance\n& Loss Intelligence"
p_main.font.name = FONT_HEADING
p_main.font.size = Pt(36)
p_main.font.bold = True
p_main.font.color.rgb = COLOR_TEXT_PRIMARY

# Subtitle
sub_box = s1.shapes.add_textbox(Inches(1.0), Inches(3.4), Inches(11.3), Inches(0.8))
tf_sub = sub_box.text_frame
tf_sub.word_wrap = True
p_sub = tf_sub.paragraphs[0]
p_sub.text = "Inverter-Level Baselining, Peer-Cohort Benchmarking, and Operational Shortfall Quantification"
p_sub.font.name = FONT_BODY
p_sub.font.size = Pt(16)
p_sub.font.color.rgb = COLOR_CYAN

# Scope Cards
card_w = Inches(3.55)
card_h = Inches(1.8)
card_y = Inches(4.4)

# Scope Card 1: Phase I
add_card(s1, Inches(1.0), card_y, card_w, card_h)
tb_c1 = s1.shapes.add_textbox(Inches(1.15), card_y + Inches(0.15), card_w - Inches(0.3), card_h - Inches(0.3))
tf_c1 = tb_c1.text_frame
tf_c1.word_wrap = True
p1 = tf_c1.paragraphs[0]
p1.text = "PHASE I • IMPLEMENTED"
p1.font.size = Pt(10)
p1.font.bold = True
p1.font.color.rgb = COLOR_EMERALD
p2 = tf_c1.add_paragraph()
p2.text = "Commercial Fleet Intelligence"
p2.font.size = Pt(14)
p2.font.bold = True
p2.font.color.rgb = COLOR_TEXT_PRIMARY
p3 = tf_c1.add_paragraph()
p3.text = "44 Inverters, 77k daylight records, 10-feature Random Forest baseline, dual-baseline detection & persistence filtering."
p3.font.size = Pt(10.5)
p3.font.color.rgb = COLOR_TEXT_MUTED

# Scope Card 2: Phase II
add_card(s1, Inches(4.88), card_y, card_w, card_h)
tb_c2 = s1.shapes.add_textbox(Inches(5.03), card_y + Inches(0.15), card_w - Inches(0.3), card_h - Inches(0.3))
tf_c2 = tb_c2.text_frame
tf_c2.word_wrap = True
p1 = tf_c2.paragraphs[0]
p1.text = "PHASE II • IMPLEMENTED"
p1.font.size = Pt(10)
p1.font.bold = True
p1.font.color.rgb = COLOR_CYAN
p2 = tf_c2.add_paragraph()
p2.text = "Residential Solar Intelligence"
p2.font.size = Pt(14)
p2.font.bold = True
p2.font.color.rgb = COLOR_TEXT_PRIMARY
p3 = tf_c2.add_paragraph()
p3.text = "Real Havells 5.0 kW single-phase export, dynamic PII sanitization, standby gap analysis, and 3-way energy reconciliation."
p3.font.size = Pt(10.5)
p3.font.color.rgb = COLOR_TEXT_MUTED

# Scope Card 3: Phase III
add_card(s1, Inches(8.76), card_y, card_w, card_h)
tb_c3 = s1.shapes.add_textbox(Inches(8.91), card_y + Inches(0.15), card_w - Inches(0.3), card_h - Inches(0.3))
tf_c3 = tb_c3.text_frame
tf_c3.word_wrap = True
p1 = tf_c3.paragraphs[0]
p1.text = "PHASE III • PLANNED ROADMAP"
p1.font.size = Pt(10)
p1.font.bold = True
p1.font.color.rgb = COLOR_AMBER
p2 = tf_c3.add_paragraph()
p2.text = "Forecasting & Live Inverters"
p2.font.size = Pt(14)
p2.font.bold = True
p2.font.color.rgb = COLOR_TEXT_PRIMARY
p3 = tf_c3.add_paragraph()
p3.text = "NWP day-ahead solar forecasting, localized cloud modeling, and authorized Modbus/Cloud live inverter ingestion."
p3.font.size = Pt(10.5)
p3.font.color.rgb = COLOR_TEXT_MUTED

# Footer metadata
meta_box = s1.shapes.add_textbox(Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.4))
tf_m = meta_box.text_frame
p_m = tf_m.paragraphs[0]
p_m.text = "Department of Computer Science & Engineering • Academic Project Defense • 46 Automated Unit Tests Passing"
p_m.font.size = Pt(11)
p_m.font.color.rgb = COLOR_TEXT_MUTED

# =============================================================================
# SLIDE 2: PROBLEM & MOTIVATION
# =============================================================================
s2 = prs.slides.add_slide(blank_layout)
set_slide_background(s2)
add_header(s2, "The Operational Challenge", "Why Falling Solar Power Does Not Imply a Fault")

col_w = Inches(5.6)
card_y = Inches(1.5)
card_h = Inches(5.3)

# Left Column: The Problem & Traditional Pitfalls
add_card(s2, Inches(0.8), card_y, col_w, card_h)
tb_p = s2.shapes.add_textbox(Inches(1.05), card_y + Inches(0.25), col_w - Inches(0.5), card_h - Inches(0.5))
tf_p = tb_p.text_frame
tf_p.word_wrap = True

p = tf_p.paragraphs[0]
p.text = "THE PROBLEM IN SOLAR TELEMETRY"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_ROSE

p = tf_p.add_paragraph()
p.text = "Solar power generation fluctuates continuously with sun position and passing cloud cover. A drop in power is frequently normal meteorological behavior, NOT an equipment failure."
p.font.size = Pt(12)
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_p.add_paragraph()
p.text = "\nTWO CRITICAL INDUSTRY PITFALLS:"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_p.add_paragraph()
p.text = "1. The False Alarm Trap (Passing Clouds):"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_AMBER
p = tf_p.add_paragraph()
p.text = "Static SCADA thresholds (e.g., P < 0.2 kW) trigger hundreds of nuisance alerts daily during cloud shadows. In our test data, 47.6 kWh of deviation was transient cloud noise."
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_p.add_paragraph()
p.text = "2. The Invisible Degrade (Masked Underperformance):"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_AMBER
p = tf_p.add_paragraph()
p.text = "A string disconnect, severe soiling, or inverter derating (-30%) still produces 1.0 kW on sunny days. Static thresholds completely miss this, allowing thousands of kWh to be silently lost."
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_MUTED

# Right Column: The Solution
add_card(s2, Inches(6.8), card_y, col_w, card_h)
tb_s = s2.shapes.add_textbox(Inches(7.05), card_y + Inches(0.25), col_w - Inches(0.5), card_h - Inches(0.5))
tf_s = tb_s.text_frame
tf_s.word_wrap = True

p = tf_s.paragraphs[0]
p.text = "OUR ENGINEERING SOLUTION"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD

p = tf_s.add_paragraph()
p.text = "Adaptive Solar Intelligence bridges physical meteorology and machine learning to answer four fundamental operator questions:"
p.font.size = Pt(12)
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_s.add_paragraph()
p.text = "\n1. What SHOULD this inverter generate right now?"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN
p = tf_s.add_paragraph()
p.text = "Model-predicted active power baseline conditioned on pyranometer irradiance, module temperature, and solar geometry."
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_s.add_paragraph()
p.text = "2. How does it compare to its immediate neighbors?"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN
p = tf_s.add_paragraph()
p.text = "Capacity-normalized peer cohort benchmark prevents the ML model from normalizing chronic long-term derating as baseline."
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_s.add_paragraph()
p.text = "3. Has the underperformance persisted over time?"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN
p = tf_s.add_paragraph()
p.text = "Strict temporal persistence filter (>= 45 min) suppresses short cloud transients while catching sustained faults."
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_s.add_paragraph()
p.text = "4. What is the cumulative energy lost (kWh)?"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN
p = tf_s.add_paragraph()
p.text = "Quantifies potential shortfall in kilowatt-hours and ranks inverters by urgency to prioritize maintenance dispatch."
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_MUTED

# =============================================================================
# SLIDE 3: SYSTEM ARCHITECTURE
# =============================================================================
s3 = prs.slides.add_slide(blank_layout)
set_slide_background(s3)
add_header(s3, "System Architecture", "End-to-End Solar Intelligence Pipeline")

# We will create 6 sequential architectural pipeline stages as cards
stage_w = Inches(1.8)
stage_h = Inches(4.8)
stage_y = Inches(1.6)

stages = [
    ("1. INGESTION & UNITS", COLOR_CYAN, [
        ("Raw SCADA Data", "Plants 1 & 2 (4 CSVs)"),
        ("Unit Resolution", "Watts to canonical kW"),
        ("Weather Archive", "Open-Meteo ERA5 15m"),
        ("Join Invariants", "Assert zero row mult")
    ]),
    ("2. FEATURES", COLOR_CYAN, [
        ("Diurnal Encodings", "sin/cos time-of-day"),
        ("Seasonal Encodings", "sin/cos day-of-year"),
        ("Thermal Gradient", "Module - Ambient Temp"),
        ("Leakage Shield", "Exclude DC & yields")
    ]),
    ("3. VALIDATION", COLOR_CYAN, [
        ("Chronological Order", "Strict timestamp sort"),
        ("Train Set (60%)", "May 15 - Jun 4 (45k)"),
        ("Val Set (20%)", "Jun 4 - Jun 11 (16k)"),
        ("Out-of-Time Test", "Jun 11 - Jun 17 (16k)")
    ]),
    ("4. ML MODELING", COLOR_GOLD, [
        ("Regression Target", "AC_POWER_KW"),
        ("Selected Pipeline", "Random Forest (100t)"),
        ("Occam's Razor", "Baseline beats weather"),
        ("Model Recompute", "Test MAE = 70.4 W")
    ]),
    ("5. DETECTION", COLOR_AMBER, [
        ("Dual Baseline", "Inverter ML + Peer Med"),
        ("Capacity Normal", "Scale by 99th pct capacity"),
        ("Persistence Filter", ">= 3 intervals (>= 45m)"),
        ("Gap Reset", "No cross-gap bridging")
    ]),
    ("6. LOSS INTELLIGENCE", COLOR_EMERALD, [
        ("Energy Integration", "dE = P_loss x 0.25h"),
        ("Urgency Ranking", "Sorted by lost kWh"),
        ("Noise Filtered", "47.6 kWh transient rej"),
        ("Decision Support", "6-View Streamlit UI")
    ]),
]

for idx, (title, color, bullets) in enumerate(stages):
    card_x = Inches(0.8) + idx * Inches(1.97)
    add_card(s3, card_x, stage_y, stage_w, stage_h)

    tb = s3.shapes.add_textbox(card_x + Inches(0.1), stage_y + Inches(0.15), stage_w - Inches(0.2), stage_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = color

    for b_title, b_desc in bullets:
        p = tf.add_paragraph()
        p.text = f"\n{b_title}"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_PRIMARY

        p = tf.add_paragraph()
        p.text = b_desc
        p.font.size = Pt(9)
        p.font.color.rgb = COLOR_TEXT_MUTED

# Residential Validation Track Callout
add_card(s3, Inches(0.8), Inches(6.55), Inches(11.65), Inches(0.7), bg_color=COLOR_SURFACE_ALT, border_color=COLOR_CYAN)
tb_callout = s3.shapes.add_textbox(Inches(1.0), Inches(6.6), Inches(11.25), Inches(0.6))
tf_c = tb_callout.text_frame
tf_c.word_wrap = True
p = tf_c.paragraphs[0]
p.text = "DEDICATED RESIDENTIAL VALIDATION TRACK (PHASE II): "
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN
p.add_run().text = "Real Havells 5.0 kW single-phase export (997 records, 7 days) evaluated in parallel with PII stripping, standby gap analysis, and 3-way energy reconciliation. Never merged into utility training data."
p.runs[1].font.size = Pt(10.5)
p.runs[1].font.color.rgb = COLOR_TEXT_PRIMARY

# =============================================================================
# SLIDE 4: DATASET & PREPROCESSING
# =============================================================================
s4 = prs.slides.add_slide(blank_layout)
set_slide_background(s4)
add_header(s4, "Data Engineering & Physical Grounding", "Verified Empirical Schemas & Leakage Prevention")

# 3 Cards: Left = 1,000x Discrepancy, Mid = Dataset Stats, Right = Leakage Elimination
c_w = Inches(3.68)
c_y = Inches(1.5)
c_h = Inches(5.4)

# Card 1: Physical Unit Reconciliation
add_card(s4, Inches(0.8), c_y, c_w, c_h)
tb_u = s4.shapes.add_textbox(Inches(1.0), c_y + Inches(0.2), c_w - Inches(0.4), c_h - Inches(0.4))
tf_u = tb_u.text_frame
tf_u.word_wrap = True

p = tf_u.paragraphs[0]
p.text = "1,000x UNIT RESOLUTION"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_AMBER

p = tf_u.add_paragraph()
p.text = "The raw Kaggle dataset labeled AC_POWER as kW, but empirical values peaked at 1,410.95 (Plant 1) and 1,385.42 (Plant 2)."
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_u.add_paragraph()
p.text = "\nPublished Ground Truth:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_u.add_paragraph()
p.text = "A peer-reviewed benchmark (Clean Energy, Oxford Academic, 2026) verified these exact plants in Karnataka:"
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_u.add_paragraph()
p.text = "• 22 Growatt MIC 1500TL-X inverters\n• Rated capacity: 1.5 kW AC each\n• Total array capacity: 33.0 kW AC\n• Empirical peak sum: 29.15 kW"
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_CYAN

p = tf_u.add_paragraph()
p.text = "\nCanonical Unit Scaling:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD

p = tf_u.add_paragraph()
p.text = "Raw values are Watts (W), NOT kW! Established AC_POWER_KW = AC_POWER / 1000.0 across all pipelines, preventing a 1,000x distortion in loss calculations."
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_MUTED

# Card 2: Dataset Statistics
add_card(s4, Inches(4.78), c_y, c_w, c_h)
tb_d = s4.shapes.add_textbox(Inches(4.98), c_y + Inches(0.2), c_w - Inches(0.4), c_h - Inches(0.4))
tf_d = tb_d.text_frame
tf_d.word_wrap = True

p = tf_d.paragraphs[0]
p.text = "EMPIRICAL DATASET VOLUMES"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN

p = tf_d.add_paragraph()
p.text = "Commercial Utility Fleet (Kaggle):"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_d.add_paragraph()
p.text = "• Plant 1 Gen: 68,778 rows (22 inverters)\n• Plant 1 Weather: 3,182 rows (1 sensor)\n• Plant 2 Gen: 67,698 rows (22 inverters)\n• Plant 2 Weather: 3,259 rows (1 sensor)\n• Joined Daylight Rows: 77,098\n• Interval: 15-min (May 15 - Jun 17, 2020)"
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_d.add_paragraph()
p.text = "\nResidential Export (Havells 5.0 kW):"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_d.add_paragraph()
p.text = "• Telemetry Records: 997 synchronous rows\n• Telemetry Channels: 63 columns\n• Cadence: 5.0-minute (Oct 20-26, 2026)\n• Data Quality: 0 nulls, 0 duplicate timestamps\n• Standby Cycles: 6 overnight gaps (~12h)\n• Daytime Communications: 0 dropouts"
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_MUTED

# Card 3: Leakage Safeguards
add_card(s4, Inches(8.76), c_y, c_w, c_h)
tb_l = s4.shapes.add_textbox(Inches(8.96), c_y + Inches(0.2), c_w - Inches(0.4), c_h - Inches(0.4))
tf_l = tb_l.text_frame
tf_l.word_wrap = True

p = tf_l.paragraphs[0]
p.text = "STRICT LEAKAGE SAFEGUARDS"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD

p = tf_l.add_paragraph()
p.text = "Target Leakage Elimination:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_l.add_paragraph()
p.text = "DC_POWER, DAILY_YIELD, and TOTAL_YIELD were strictly excluded from model features. Including them produces a trivial R2 > 0.999 that cannot detect when AC power drops relative to weather."
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_l.add_paragraph()
p.text = "\nTemporal Transformation Boundaries:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_l.add_paragraph()
p.text = "All scalers, encoders, and threshold quantiles are fitted strictly on training data slices. Never fit or transform across split boundaries."
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_l.add_paragraph()
p.text = "\nCyclical Solar Time Encodings:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_l.add_paragraph()
p.text = "Derived sin/cos time-of-day and day-of-year features to capture smooth diurnal solar curves without discontinuous step boundaries."
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_MUTED

# =============================================================================
# SLIDE 5: AIML METHODOLOGY & MODEL SELECTION
# =============================================================================
s5 = prs.slides.add_slide(blank_layout)
set_slide_background(s5)
add_header(s5, "AIML Methodology & Model Selection", "Chronological Out-of-Time Benchmark")

# Left Column: Task & Features (Width: 4.8)
add_card(s5, Inches(0.8), Inches(1.5), Inches(4.8), Inches(5.4))
tb_m = s5.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(4.4), Inches(5.0))
tf_m = tb_m.text_frame
tf_m.word_wrap = True

p = tf_m.paragraphs[0]
p.text = "THE REGRESSION TASK"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN

p = tf_m.add_paragraph()
p.text = "Predict expected continuous active power P_expected (kW) conditioned on 10 exogenous meteorological and cyclical time features:"
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_m.add_paragraph()
p.text = "\n10 Leakage-Free Input Features:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_AMBER

p = tf_m.add_paragraph()
p.text = "1. IRRADIATION (pyranometer kW/m2)\n2. AMBIENT_TEMPERATURE (°C)\n3. MODULE_TEMPERATURE (°C)\n4. TEMP_DIFF (Module - Ambient °C)\n5. sin_time_of_day (diurnal cycle)\n6. cos_time_of_day (diurnal cycle)\n7. sin_day_of_year (seasonal cycle)\n8. cos_day_of_year (seasonal cycle)\n9. PLANT_ID (facility embedding)\n10. SOURCE_KEY (inverter identifier)"
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_m.add_paragraph()
p.text = "\nStrict Chronological Partitioning:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD

p = tf_m.add_paragraph()
p.text = "• Train (60%): May 15 - Jun 4 (44,966 rows)\n• Val (20%): Jun 4 - Jun 11 (15,896 rows)\n• Test (20%): Jun 11 - Jun 17 (16,236 rows)\nNever shuffle time-series data!"
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_MUTED

# Right Column: Model Selection Benchmark Table (Width: 6.6)
add_card(s5, Inches(5.9), Inches(1.5), Inches(6.6), Inches(5.4))
tb_t = s5.shapes.add_textbox(Inches(6.1), Inches(1.7), Inches(6.2), Inches(5.0))
tf_t = tb_t.text_frame
tf_t.word_wrap = True

p = tf_t.paragraphs[0]
p.text = "OUT-OF-TIME BENCHMARK (TEST WEEK: JUNE 11-17, 2020)"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_GOLD

p = tf_t.add_paragraph()
p.text = "Evaluated on 16,236 unseen chronological daylight observations:"
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_MUTED

# Add formatted table inside slide
table_shape = s5.shapes.add_table(7, 6, Inches(6.1), Inches(2.3), Inches(6.2), Inches(2.6))
table = table_shape.table

# Set column widths
table.columns[0].width = Inches(2.2)
table.columns[1].width = Inches(0.8)
table.columns[2].width = Inches(0.8)
table.columns[3].width = Inches(0.8)
table.columns[4].width = Inches(0.8)
table.columns[5].width = Inches(0.8)

headers = ["Model Pipeline", "Val MAE", "Test MAE", "Test RMSE", "Test R²", "Test WAPE"]
rows = [
    ["Baseline (RandomForest)*", "0.1055 kW", "0.0704 kW", "0.1704 kW", "0.7337", "16.46%"],
    ["Enriched (RandomForest)", "0.1076 kW", "0.0708 kW", "0.1660 kW", "0.7473", "16.56%"],
    ["Baseline (HistGBR)", "0.1073 kW", "0.0699 kW", "0.1646 kW", "0.7515", "16.35%"],
    ["Enriched (HistGBR)", "0.1105 kW", "0.0735 kW", "0.1634 kW", "0.7550", "17.19%"],
    ["Baseline (Ridge)", "0.1257 kW", "0.1120 kW", "0.1814 kW", "0.6982", "26.18%"],
    ["Enriched (Ridge)", "0.1287 kW", "0.1158 kW", "0.1817 kW", "0.6972", "27.08%"],
]

for col_idx, h in enumerate(headers):
    cell = table.cell(0, col_idx)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_SURFACE_ALT
    p = cell.text_frame.paragraphs[0]
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

for row_idx, r_data in enumerate(rows):
    is_winner = (row_idx == 0)
    for col_idx, val in enumerate(r_data):
        cell = table.cell(row_idx + 1, col_idx)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(26, 44, 40) if is_winner else COLOR_SURFACE
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(9)
        p.font.bold = is_winner
        p.font.color.rgb = COLOR_EMERALD if is_winner else COLOR_TEXT_PRIMARY

# Add caption below table
tb_cap = s5.shapes.add_textbox(Inches(6.1), Inches(5.1), Inches(6.2), Inches(1.6))
tf_cap = tb_cap.text_frame
tf_cap.word_wrap = True
p = tf_cap.paragraphs[0]
p.text = "*Selection Rationale (Baseline RandomForest Retained):"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD
p = tf_cap.add_paragraph()
p.text = "1. Lowest Validation Error (0.1055 kW MAE): Model selection strictly followed validation ranking without looking at test labels.\n2. Recomputed Test MAE of 70.4 W is just 4.7% of inverter nameplate capacity.\n3. R² of 0.7337 reflects strong predictive fidelity under highly dynamic solar conditions."
p.font.size = Pt(9.5)
p.font.color.rgb = COLOR_TEXT_MUTED

# =============================================================================
# SLIDE 6: EXPERIMENTAL RESULTS & WEATHER ENRICHMENT
# =============================================================================
s6 = prs.slides.add_slide(blank_layout)
set_slide_background(s6)
add_header(s6, "Experimental Analysis", "The Regional Weather Reanalysis Paradox")

col_w = Inches(5.6)
card_y = Inches(1.5)
card_h = Inches(5.4)

# Left Column: Weather Experiment
add_card(s6, Inches(0.8), card_y, col_w, card_h)
tb_w = s6.shapes.add_textbox(Inches(1.05), card_y + Inches(0.25), col_w - Inches(0.5), card_h - Inches(0.5))
tf_w = tb_w.text_frame
tf_w.word_wrap = True

p = tf_w.paragraphs[0]
p.text = "THE WEATHER ENRICHMENT EXPERIMENT"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN

p = tf_w.add_paragraph()
p.text = "We evaluated whether enriching onsite telemetry with regional ERA5 reanalysis weather from Open-Meteo (temperature, humidity, cloud cover, wind speed) improves power prediction."
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_w.add_paragraph()
p.text = "\nEmpirical Result:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_AMBER

p = tf_w.add_paragraph()
p.text = "Adding regional weather slightly DEGRADED validation error (+1.99% MAE, from 0.1055 to 0.1076 kW). It did NOT improve accuracy."
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_w.add_paragraph()
p.text = "\nPhysical & Scientific Explanation:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD

p = tf_w.add_paragraph()
p.text = "1. Spatial Resolution Disparity:\nERA5 reanalysis operates on a 0.1° (~11 km) spatial grid. It cannot resolve localized cloud shadows passing over individual solar arrays."
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_w.add_paragraph()
p.text = "2. Pyranometer Superiority:\nOnsite pyranometers and module thermocouples capture direct solar physics and cell temperature at the exact array location with zero spatial lag."
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_w.add_paragraph()
p.text = "3. Occam's Razor Selection:\nThe simpler 10-feature baseline was retained: lower error, lower complexity, and zero external API dependencies."
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_MUTED

# Right Column: Honest Metric Defense
add_card(s6, Inches(6.8), card_y, col_w, card_h)
tb_e = s6.shapes.add_textbox(Inches(7.05), card_y + Inches(0.25), col_w - Inches(0.5), card_h - Inches(0.5))
tf_e = tb_e.text_frame
tf_e.word_wrap = True

p = tf_e.paragraphs[0]
p.text = "HONEST METRIC INTERPRETATION"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_GOLD

p = tf_e.add_paragraph()
p.text = "Recomputed Metrics from Saved Artifact (models/best_solar_model.joblib):"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY

# Metric Pill Rows
metrics = [
    ("MAE = 0.0704 kW (70.4 W)", "Mean Absolute Error: Average prediction uncertainty is just 70.4 Watts across 1,500 Watt inverters (4.7% of capacity)."),
    ("RMSE = 0.1704 kW (170.4 W)", "Root Mean Squared Error: Penalizes larger deviations during steep morning irradiance ramps."),
    ("R² = 0.7337", "Coefficient of Determination: Explains 73.4% of power variance out-of-time. We NEVER call R² 'accuracy'."),
    ("WAPE = 16.46%", "Weighted Absolute Percentage Error: Scale-independent percentage error across the entire daylight fleet.")
]

for m_title, m_desc in metrics:
    p = tf_e.add_paragraph()
    p.text = f"\n• {m_title}"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p = tf_e.add_paragraph()
    p.text = f"  {m_desc}"
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_e.add_paragraph()
p.text = "\nZero Data Leakage Assertion:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD
p = tf_e.add_paragraph()
p.text = "Verified: No DC_POWER, DAILY_YIELD, or future timestamps were accessible to the model. Out-of-time evaluation guarantees defensibility."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

# =============================================================================
# SLIDE 7: UNDERPERFORMANCE & LOSS INTELLIGENCE
# =============================================================================
s7 = prs.slides.add_slide(blank_layout)
set_slide_background(s7)
add_header(s7, "Underperformance & Loss Intelligence", "Dual Baseline, Persistence Filtering & Energy Accounting")

# Top Row: Mathematical Foundations (Left: Dual Baseline, Right: Persistence)
card_top_w = Inches(5.6)
card_top_h = Inches(2.6)

# Card 1: Dual Baseline
add_card(s7, Inches(0.8), Inches(1.5), card_top_w, card_top_h)
tb_db = s7.shapes.add_textbox(Inches(1.0), Inches(1.65), card_top_w - Inches(0.4), card_top_h - Inches(0.3))
tf_db = tb_db.text_frame
tf_db.word_wrap = True
p = tf_db.paragraphs[0]
p.text = "DUAL-BASELINE DETECTION LOGIC"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN

p = tf_db.add_paragraph()
p.text = "1. Inverter ML Baseline: Residual = P_expected - P_actual"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY
p = tf_db.add_paragraph()
p.text = "Catches sudden drops relative to weather. Positive residual = deficit."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_db.add_paragraph()
p.text = "2. Peer-Cohort Baseline: P_cohort = Median({P_j / C_j}) x C_i"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY
p = tf_db.add_paragraph()
p.text = "Normalized by capacity C_i. Prevents the ML model from normalizing chronic underperformance as expected behavior."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

# Card 2: Persistence & Thresholds
add_card(s7, Inches(6.8), Inches(1.5), card_top_w, card_top_h)
tb_pf = s7.shapes.add_textbox(Inches(7.0), Inches(1.65), card_top_w - Inches(0.4), card_top_h - Inches(0.3))
tf_pf = tb_pf.text_frame
tf_pf.word_wrap = True
p = tf_pf.paragraphs[0]
p.text = "PERSISTENCE FILTER & THRESHOLD CALIBRATION"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = COLOR_AMBER

p = tf_pf.add_paragraph()
p.text = ">= 3 Consecutive Intervals (>= 45 Minutes):"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY
p = tf_pf.add_paragraph()
p.text = "Anomalies must persist across at least 3 intervals to trigger alerts. Missing timestamps immediately reset the counter."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_pf.add_paragraph()
p.text = "Outage-Filtered Threshold Recalibration:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY
p = tf_pf.add_paragraph()
p.text = "Plant 2 high-sun threshold recalibrated to 0.0846 kW on operating inverters, preventing outages from masking subtle derates."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

# Bottom Row: Operational Shortfall Results (4 Metric Cards)
kpi_w = Inches(2.7)
kpi_h = Inches(2.4)
kpi_y = Inches(4.35)

kpis = [
    ("1,736.3 kWh", "Measured Generation", "Actual energy generated by 44 inverters during out-of-time test week.", COLOR_TEXT_PRIMARY),
    ("1,724.5 kWh", "Expected Generation", "Baseline generation predicted by 10-feature Random Forest model.", COLOR_CYAN),
    ("89.44 kWh (5.2%)", "Persistent Shortfall", "Actionable energy loss lasting >= 45m requiring maintenance dispatch.", COLOR_AMBER),
    ("47.58 kWh", "Transient Noise Filtered", "Passing cloud noise safely suppressed by persistence filter (34.7% of raw dev).", COLOR_EMERALD),
]

for idx, (val, title, desc, color) in enumerate(kpis):
    kx = Inches(0.8) + idx * Inches(2.98)
    add_card(s7, kx, kpi_y, kpi_w, kpi_h)

    tb = s7.shapes.add_textbox(kx + Inches(0.15), kpi_y + Inches(0.2), kpi_w - Inches(0.3), kpi_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = val
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = color

    p = tf.add_paragraph()
    p.text = title
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PRIMARY

    p = tf.add_paragraph()
    p.text = desc
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_TEXT_MUTED

# =============================================================================
# SLIDE 8: WORKING PROTOTYPE DEMONSTRATION
# =============================================================================
s8 = prs.slides.add_slide(blank_layout)
set_slide_background(s8)
add_header(s8, "Working Prototype Demonstration", "Interactive Streamlit Operations Dashboard (6 Views)")

# 4 Image Cards in a 2x2 grid displaying actual verified screenshots
img_w = Inches(5.6)
img_h = Inches(2.45)

grid_items = [
    ("Overview & Operational Loss Summary", FIGURES_DIR / "audit_view_1_overview.png", Inches(0.8), Inches(1.5)),
    ("Plant & Inverter Telemetry Explorer", FIGURES_DIR / "audit_view_3_explorer.png", Inches(6.8), Inches(1.5)),
    ("Inverter Loss Prioritization Ranking", FIGURES_DIR / "audit_view_4_prioritization.png", Inches(0.8), Inches(4.25)),
    ("Interactive Synthetic Detection Sandbox", FIGURES_DIR / "audit_view_5_synthetic.png", Inches(6.8), Inches(4.25)),
]

for title, path, gx, gy in grid_items:
    add_card(s8, gx, gy, img_w, img_h)

    # Title overlay bar
    tb = s8.shapes.add_textbox(gx + Inches(0.15), gy + Inches(0.08), img_w - Inches(0.3), Inches(0.35))
    tf = tb.text_frame
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

    # Embed image if exists
    if path.exists():
        s8.shapes.add_picture(str(path), gx + Inches(0.15), gy + Inches(0.42), img_w - Inches(0.3), img_h - Inches(0.52))

# =============================================================================
# SLIDE 9: PHASE II RESIDENTIAL SOLAR INTELLIGENCE
# =============================================================================
s9 = prs.slides.add_slide(blank_layout)
set_slide_background(s9)
add_header(s9, "Phase II Residential Validation", "Real Havells 5.0 kW Single-Phase Inverter Audit")

col_w = Inches(5.6)
card_y = Inches(1.5)
card_h = Inches(5.4)

# Left Column: Telemetry Audit & 3-Way Reconciliation
add_card(s9, Inches(0.8), card_y, col_w, card_h)
tb_r = s9.shapes.add_textbox(Inches(1.05), card_y + Inches(0.2), col_w - Inches(0.5), card_h - Inches(0.4))
tf_r = tb_r.text_frame
tf_r.word_wrap = True

p = tf_r.paragraphs[0]
p.text = "EMPIRICAL RESIDENTIAL TELEMETRY AUDIT"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN

p = tf_r.add_paragraph()
p.text = "• Inverter: Havells 5.0 kW single-phase grid-tied string inverter\n• Telemetry Volume: 997 synchronous readings (Oct 20-26, 2026)\n• Sampling Cadence: Nominal 5.0-minute (0 nulls, 0 duplicates)\n• Observed Peak: 3.58 kW (71.6% of rated 5.0 kW capacity)\n• Standby Detection: 6 nocturnal sleep gaps (~12h); 0 daytime dropouts\n• Privacy Protection: Serial numbers and account IDs scrubbed dynamically"
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_r.add_paragraph()
p.text = "\nTHREE-WAY ENERGY RECONCILIATION:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD

p = tf_r.add_paragraph()
p.text = "1. Cumulative Odometer Delta: 7,992.0 - 7,857.1 = 134.90 kWh\n2. Daily Reset Counter Sum: Sum(Daily Yield) = 134.94 kWh\n3. Discrete Trapezoidal Integral: Sum(P x dt) = 135.90 kWh (+0.7%)"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_r.add_paragraph()
p.text = "\nOperating State Distribution:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_AMBER

p = tf_r.add_paragraph()
p.text = "Grid Connected (96.8%, 965), Waiting (2.3%, 23 at dawn/dusk), Recoverable Fault (0.7%, 7 grid excursions), Checking (0.2%, 2)."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

# Right Column: Screenshot & Honest Roadmap
add_card(s9, Inches(6.8), card_y, col_w, card_h)

# Embed Phase II Screenshot
p2_img = FIGURES_DIR / "audit_view_2_phase_ii.png"
if p2_img.exists():
    s9.shapes.add_picture(str(p2_img), Inches(7.0), card_y + Inches(0.2), col_w - Inches(0.4), Inches(2.9))

tb_f = s9.shapes.add_textbox(Inches(7.0), card_y + Inches(3.2), col_w - Inches(0.4), Inches(2.0))
tf_f = tb_f.text_frame
tf_f.word_wrap = True

p = tf_f.paragraphs[0]
p.text = "HONEST FORECAST READINESS SCORECARD"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_GOLD

p = tf_f.add_paragraph()
p.text = "• Telemetry Pipeline: READY (Verified 7-day ingestion & sanitization)\n• External Weather: PENDING ALIGNMENT (Requires verified rooftop GPS)\n• Live Inverter API: PENDING AUTHORIZATION (Modbus TCP/Cloud webhook)\n• Validated Day-Ahead Model: PHASE III ROADMAP"
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_f.add_paragraph()
p.text = "\nDefense Boundary: A 7-day export is valuable for baseline profiling, but CANNOT train a multi-season day-ahead forecast without extreme seasonal overfitting. We present no fabricated predictions."
p.font.size = Pt(9.5)
p.font.color.rgb = COLOR_AMBER

# =============================================================================
# SLIDE 10: CONCLUSION, ROADMAP & QUESTIONS
# =============================================================================
s10 = prs.slides.add_slide(blank_layout)
set_slide_background(s10)
add_header(s10, "Summary & Future Roadmap", "Conclusions & Viva Defense Readiness")

c_w10 = Inches(3.68)
c_y10 = Inches(1.5)
c_h10 = Inches(5.4)

# Card 1: Technical Accomplishments
add_card(s10, Inches(0.8), c_y10, c_w10, c_h10)
tb_c1 = s10.shapes.add_textbox(Inches(1.0), c_y10 + Inches(0.2), c_w10 - Inches(0.4), c_h10 - Inches(0.4))
tf_c1 = tb_c1.text_frame
tf_c1.word_wrap = True

p = tf_c1.paragraphs[0]
p.text = "KEY ACCOMPLISHMENTS"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD

p = tf_c1.add_paragraph()
p.text = "1. Solved 1,000x Unit Error:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY
p = tf_c1.add_paragraph()
p.text = "Reconciled Kaggle SCADA against Oxford academic literature (1.5 kW MIC 1500TL-X inverters)."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_c1.add_paragraph()
p.text = "2. Leakage-Free ML Modeling:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY
p = tf_c1.add_paragraph()
p.text = "10-feature Random Forest baseline achieves Test MAE = 70.4 W, R2 = 0.7337 on strictly chronological splits."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_c1.add_paragraph()
p.text = "3. Actionable Loss Quantification:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY
p = tf_c1.add_paragraph()
p.text = "Suppressed 47.6 kWh of transient cloud noise while isolating 89.4 kWh of persistent actionable shortfall."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_c1.add_paragraph()
p.text = "4. Robust Working Platform:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY
p = tf_c1.add_paragraph()
p.text = "6-view Apple-inspired dashboard backed by 46 passing automated tests (0 warnings)."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

# Card 2: Explicit Limitations
add_card(s10, Inches(4.78), c_y10, c_w10, c_h10)
tb_c2 = s10.shapes.add_textbox(Inches(4.98), c_y10 + Inches(0.2), c_w10 - Inches(0.4), c_h10 - Inches(0.4))
tf_c2 = tb_c2.text_frame
tf_c2.word_wrap = True

p = tf_c2.paragraphs[0]
p.text = "EXPLICIT LIMITATIONS"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_AMBER

p = tf_c2.add_paragraph()
p.text = "1. Short Observation Horizon:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY
p = tf_c2.add_paragraph()
p.text = "Utility dataset spans 34 days (early summer); residential spans 7 days. Annualizing losses or claiming monsoon invariance is invalid."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_c2.add_paragraph()
p.text = "2. Candidate vs Confirmed Faults:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY
p = tf_c2.add_paragraph()
p.text = "Without physical maintenance work orders, detections are candidate underperformance intervals, not verified component failures."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_c2.add_paragraph()
p.text = "3. Reanalysis Weather Limits:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY
p = tf_c2.add_paragraph()
p.text = "Regional reanalysis (~11 km) cannot match the microclimate fidelity of onsite pyranometers."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

# Card 3: Roadmap & Defense
add_card(s10, Inches(8.76), c_y10, c_w10, c_h10)
tb_c3 = s10.shapes.add_textbox(Inches(8.96), c_y10 + Inches(0.2), c_w10 - Inches(0.4), c_h10 - Inches(0.4))
tf_c3 = tb_c3.text_frame
tf_c3.word_wrap = True

p = tf_c3.paragraphs[0]
p.text = "PHASE III ROADMAP & VIVA"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN

p = tf_c3.add_paragraph()
p.text = "Planned Phase III Work:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY
p = tf_c3.add_paragraph()
p.text = "• Weather-Informed Day-Ahead Forecasts: High-resolution GHI forecasts coupled to trained regression models.\n• Authorized Live Inverter Feeds: Modbus TCP gateway and cloud OAuth webhooks for real-time telemetry."
p.font.size = Pt(10)
p.font.color.rgb = COLOR_TEXT_MUTED

p = tf_c3.add_paragraph()
p.text = "\nEvaluation Readiness:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_GOLD
p = tf_c3.add_paragraph()
p.text = "• 46 Automated Tests Passing (100%)\n• Recomputed & Verified Artifacts\n• Interactive 6-View Streamlit App\n• Open for Questions & Demonstration"
p.font.size = Pt(10.5)
p.font.color.rgb = COLOR_TEXT_PRIMARY

p = tf_c3.add_paragraph()
p.text = "\nThank you. We welcome your questions."
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN

# Save Presentation
output_path = PRESENTATION_DIR / "Adaptive_Solar_Performance_Loss_Intelligence_Evaluation.pptx"
prs.save(str(output_path))
print(f"SUCCESS: Generated presentation at {output_path} ({len(prs.slides)} slides)")
